"""Workbook writer: append MVAI result columns to the operator's source file.

The export contract (:data:`app.config.constants.EXPORT_EXTRA_COLUMNS`)
promises that every column the platform adds is *prefixed* with ``MVAI`` so it
can never collide with the operator's own headers, and that the source
workbook itself is what comes back - title blocks, extra sheets, formatting
and all.

Two paths exist:

* **Source available** (the normal case - uploads are retained per batch):
  the original workbook is loaded, the header row is located with the same
  alias logic the parser uses, the ten ``MVAI ...`` columns are appended after
  the last used column, and each material's results are written on its own
  source row (``Material.row_number``), so the export lines up with what the
  operator sees.
* **Source missing** (uploads purged): a fresh single-sheet workbook is built
  from the canonical fields plus the MVAI columns.  Different layout, same
  information - the export never fails just because the original file is gone.

Writing is pure in-memory (``BytesIO``); persistence and ``ExportHistory``
belong to the export service.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.config import get_settings
from app.config.constants import EXPORT_EXTRA_COLUMNS, EXPORT_FIELDS
from app.config.logging import get_logger
from app.services.excel_parser import SpreadsheetError
from app.utils.text import normalise_header

logger = get_logger(__name__)


@dataclass(slots=True)
class ExportRow:
    """The results written for one material, keyed by its source row number."""

    #: 1-based row number in the source worksheet.
    row_number: int
    material_code: str | None = None
    description: str = ""
    brand: str | None = None
    model: str | None = None
    category: str | None = None
    unit: str | None = None
    quantity: float | None = None
    location: str | None = None
    remarks: str | None = None
    status: str = ""
    image_file: str | None = None
    image_count: int = 0
    confidence: float | None = None
    image_source: str | None = None
    source_page: str | None = None
    search_query: str | None = None
    last_error: str | None = None
    extra: dict[str, Any] = field(default_factory=dict)

    def export_values(self) -> dict[str, Any]:
        """Canonical-field view used by the fallback workbook and CSV/JSON."""
        return {
            "row_number": self.row_number,
            "material_code": self.material_code,
            "description": self.description,
            "brand": self.brand,
            "model": self.model,
            "category": self.category,
            "unit": self.unit,
            "quantity": self.quantity,
            "location": self.location,
            "remarks": self.remarks,
            "status": self.status,
            "confidence": self.confidence,
            "image_count": self.image_count,
            "accepted_image_path": self.image_file,
            "accepted_image_source": self.image_source,
            "accepted_image_page": self.source_page,
            "search_query": self.search_query,
            "last_error": self.last_error,
        }

    def mvai_values(self) -> tuple[Any, ...]:
        """Values for :data:`EXPORT_EXTRA_COLUMNS`, in that exact order."""
        confidence = (
            f"{self.confidence:.2f}" if self.confidence is not None else None
        )
        return (
            self.status,
            self.image_file,
            self.image_count,
            confidence,
            self.image_source,
            self.source_page,
            self.brand,
            self.model,
            self.search_query,
            self.last_error,
        )


class WorkbookWriter:
    """Appends MVAI result columns to a workbook (source or rebuilt)."""

    def __init__(self, *, header_scan_rows: int | None = None) -> None:
        settings = get_settings()
        self.header_scan_rows = header_scan_rows or settings.excel_header_scan_rows
        self._aliases = settings.column_aliases

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------
    def write(self, source: bytes | None, rows: Iterable[ExportRow]) -> bytes:
        """Return the export workbook as ``.xlsx`` bytes.

        Args:
            source: the operator's original workbook, or ``None`` to build a
                fresh sheet from the canonical fields.
            rows: one :class:`ExportRow` per material, keyed by source row.
        """
        row_list = list(rows)
        if source:
            try:
                return self._augment(source, row_list)
            except SpreadsheetError as exc:
                logger.warning("export.source_unusable", error=str(exc))
        return self._rebuild(row_list)

    # ------------------------------------------------------------------
    # Source-augmented path
    # ------------------------------------------------------------------
    def _augment(self, source: bytes, rows: list[ExportRow]) -> bytes:
        workbook = load_workbook(filename=BytesIO(source))
        try:
            worksheet = workbook[workbook.sheetnames[0]]
            header_row, last_column = self._locate_header(worksheet)
            if header_row is None:
                raise SpreadsheetError("no header row located in source workbook")

            # Append the MVAI columns after the last used column, once.
            start_column = last_column + 1
            for offset, title in enumerate(EXPORT_EXTRA_COLUMNS):
                worksheet.cell(row=header_row, column=start_column + offset, value=title)

            by_row = {row.row_number: row for row in rows}
            for sheet_row in range(header_row + 1, worksheet.max_row + 1):
                export_row = by_row.get(sheet_row)
                if export_row is None:
                    continue
                for offset, value in enumerate(export_row.mvai_values()):
                    worksheet.cell(row=sheet_row, column=start_column + offset, value=value)

            buffer = BytesIO()
            workbook.save(buffer)
            return buffer.getvalue()
        finally:
            workbook.close()

    def _locate_header(self, worksheet: Worksheet) -> tuple[int | None, int]:
        """ ``(header_row, last_used_column)`` for the canonical header.``"""
        reverse: dict[str, str] = {}
        for field_name, aliases in self._aliases.items():
            for alias in aliases:
                reverse.setdefault(normalise_header(alias), field_name)

        last_column = 0
        for row_index in range(1, min(self.header_scan_rows, worksheet.max_row) + 1):
            values = [cell.value for cell in worksheet[row_index]]
            last_column = max(last_column, len(values))
            mapped = {
                reverse[key]
                for value in values
                if value is not None and (key := normalise_header(value)) in reverse
            }
            if "description" in mapped:
                return row_index, len(values)
        return None, last_column or 1

    # ------------------------------------------------------------------
    # Rebuild path (source workbook unavailable)
    # ------------------------------------------------------------------
    def _rebuild(self, rows: list[ExportRow]) -> bytes:
        workbook = Workbook()
        worksheet: Worksheet = workbook.active
        worksheet.title = "MVAI Export"
        headers = [name.replace("_", " ").title() for name in EXPORT_FIELDS]
        worksheet.append(headers + list(EXPORT_EXTRA_COLUMNS))
        for row in rows:
            values = row.export_values()
            worksheet.append(
                [values[name] for name in EXPORT_FIELDS] + list(row.mvai_values())
            )
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()


__all__ = ["ExportRow", "WorkbookWriter"]
