import openpyxl
from openpyxl.drawing.image import Image as XLImage
import io
import asyncio
import httpx
import os

input_file = "c:/Users/uttam.singh1/Desktop/Material Vision AI (MVAI)/backend/output_ISSMARTU_Material_Master_Template.xlsx"
output_file = "c:/Users/uttam.singh1/Desktop/FINAL_Material_Master_Template.xlsx"

async def main():
    print("Loading workbook...")
    wb = openpyxl.load_workbook(input_file)
    ws = wb.active
    
    semaphore = asyncio.Semaphore(50)
    
    async def download_image(row, url):
        if not url or not str(url).startswith("http"):
            return None
        async with semaphore:
            try:
                async with httpx.AsyncClient(timeout=10, follow_redirects=True) as client:
                    resp = await client.get(url, headers={"User-Agent": "Mozilla/5.0"})
                    if resp.status_code == 200:
                        return row, resp.content
            except Exception:
                return None
        return None

    tasks = []
    print("Extracting URLs...")
    for row in range(2, ws.max_row + 1):
        url = ws.cell(row=row, column=9).value
        tasks.append(download_image(row, url))
        
    print(f"Downloading {len(tasks)} images concurrently...")
    
    completed = 0
    for task in asyncio.as_completed(tasks):
        result = await task
        completed += 1
        if completed % 100 == 0:
            print(f"Downloaded {completed}/{len(tasks)}...")
            
        if result:
            row, img_bytes = result
            try:
                from PIL import Image
                import mimetypes
                if True in mimetypes.types_map: mimetypes.types_map[True]['.mpo'] = 'image/jpeg'
                img = Image.open(io.BytesIO(img_bytes))
                if img.mode != 'RGB':
                    img = img.convert('RGB')
                out_stream = io.BytesIO()
                img.save(out_stream, format='JPEG')
                out_stream.seek(0)
                
                xl_img = XLImage(out_stream)
                xl_img.width = 80
                xl_img.height = 80
                cell_addr = f"G{row}"
                ws.add_image(xl_img, cell_addr)
                ws.row_dimensions[row].height = 65
            except Exception as e:
                pass
                
    print("Deleting old URL column...")
    ws.delete_cols(9)
    
    print(f"Saving to {output_file}...")
    wb.save(output_file)
    print("Done!")

if __name__ == "__main__":
    asyncio.run(main())

